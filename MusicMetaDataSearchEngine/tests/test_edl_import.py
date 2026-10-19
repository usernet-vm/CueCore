"""Cross-Check gegen songs.db + Excel-Export."""
from openpyxl import load_workbook

from core import songsdb
from core.edl import EDLEvent
from core.edl_import import (cross_check, export_csv, export_excel, export_pdf, merge_cards, strip_stem, _export_rows, _soft_break, EXPORT_COLUMNS)


def _ev(event_num, clip_name, rec_in="00:00:00:00", rec_out="00:00:05:00", track="AA"):
    return EDLEvent(event_num, "AX", track, "C", rec_in, rec_out, rec_in, rec_out,
                    from_clip_name=clip_name)


def _db(tmp_path):
    return songsdb.open(tmp_path / "songs.db")


def test_cross_check_no_clip_name(tmp_path):
    con = _db(tmp_path)
    results = cross_check([_ev("001", "")], con)
    assert results[0]["status"] == "no_clip_name"


def test_cross_check_dedupes_same_clip_across_tracks(tmp_path):
    con = _db(tmp_path)
    songsdb.upsert(con, interpret="AURORA", titel="Runaway", labelname="Decca", labelcode="LC00171")
    con.commit()
    # Derselbe Clip auf Video + zwei Audiospuren, gleiche Record-Position = EIN Cue.
    same = [_ev("001", "AURORA - Runaway.wav", track="V"),
            _ev("002", "AURORA - Runaway.wav", track="A"),
            _ev("003", "AURORA - Runaway.wav", track="A2")]
    results = cross_check(same, con)
    assert len(results) == 1 and results[0]["status"] == "found_in_db"


def test_cross_check_dedupe_can_be_disabled(tmp_path):
    con = _db(tmp_path)
    same = [_ev("001", "X - Y.wav", track="V"), _ev("002", "X - Y.wav", track="A")]
    assert len(cross_check(same, con, dedupe=False)) == 2


def test_cross_check_unparsable_clip_name(tmp_path):
    con = _db(tmp_path)
    results = cross_check([_ev("001", "justonewordnoartist.mp3")], con)
    assert results[0]["status"] == "unparsable_clip_name"


def test_cross_check_found_in_db(tmp_path):
    con = _db(tmp_path)
    songsdb.upsert(con, interpret="Alanis Morissette", titel="Thank You",
                   labelname="MAVERICK", labelcode="LC01485")
    con.commit()

    results = cross_check([_ev("001", "Alanis Morissette - Thank You.mp3")], con)
    assert results[0]["status"] == "found_in_db"
    assert results[0]["labelcode"] == "LC01485"


def test_cross_check_missing_without_resolver(tmp_path):
    con = _db(tmp_path)
    results = cross_check([_ev("001", "New Artist - New Song.mp3")], con)
    assert results[0]["status"] == "missing_no_resolver"


def test_cross_check_resolves_and_writes_flywheel_entry(tmp_path):
    con = _db(tmp_path)

    def resolve_fn(artist, title):
        return {"labelname": "REPUBLIC", "labelcode": "LC07552", "komponist": "X", "texter": "Y"}

    results = cross_check([_ev("001", "New Artist - New Song.mp3", rec_out="00:00:03:12")],
                          con, resolve_fn=resolve_fn)
    assert results[0]["status"] == "resolved_and_written"
    assert results[0]["duration_sec"] == 4  # ceil(3.48) = 4

    hit = songsdb.find(con, "New Artist", "New Song")
    assert hit["labelcode"] == "LC07552"
    assert hit["source_type"] == "edl_import"
    assert hit["duration_sec"] is None      # Cue-Dauer gehört zum Projekt, nicht zum Song
    # Autoren aus dem Netz sind nur Vorschlag: auf der Karte, markiert,
    # nicht in der songs.db (eine Quelle traf sonst ein gleichnamiges fremdes Werk und galt als verifiziert)
    assert not hit["komponist"] and not hit["texter"]
    assert results[0]["komponist"] == "X" and results[0]["author_new"] is True


def test_cross_check_cue_duration_comes_from_edl_not_from_songs_db(tmp_path):
    """Die GEMA-Dauer eines Cues kommt immer aus der Schnittliste. Eine Dauer in der songs.db (volle
    Tracklänge einer Lieferung, Cue-Dauer eines früheren Projekts) darf sie nicht überschreiben."""
    con = _db(tmp_path)
    songsdb.import_library_tracks(con, [{"file": "CEZ4030_1.mp3", "interpret": "Pierre-André Athané", "titel": "A Sunny Day",
                                         "labelname": "CMG", "labelcode": "LC10347", "duration_sec": 125}],
                                  library="cezame", date_verified="2026-09-20")
    songsdb.upsert(con, interpret="AURORA", titel="Runaway", labelname="Decca", labelcode="LC00171",
                   source_type="edl_import", duration_sec=4)
    con.commit()
    results = cross_check([_ev("001", "CEZ4030_1.mp3", rec_out="00:00:30:00"),
                           _ev("002", "AURORA - Runaway.wav", rec_in="00:00:30:00", rec_out="00:01:30:00")], con)
    assert [c["status"] for c in results] == ["found_in_db", "found_in_db"]
    assert [c["duration_sec"] for c in results] == [30, 60]
    assert results[0]["file_duration_sec"] == 125              # Tracklänge der Lieferung bleibt eigenes Feld


def test_flywheel_does_not_bring_back_an_old_cue_duration(tmp_path):
    """Ein online aufgelöster 7-s-Cue darf seine Dauer nicht als Song-Eigenschaft hinterlassen,
    die im nächsten Projekt (derselbe Song, 60 s) wieder als Cue-Dauer auftaucht."""
    con = _db(tmp_path)
    first = cross_check([_ev("001", "New Artist - New Song.mp3", rec_out="00:00:07:00")], con,
                        resolve_fn=lambda a, t: {"labelname": "REPUBLIC", "labelcode": "LC07552"})
    assert first[0]["status"] == "resolved_and_written" and first[0]["duration_sec"] == 7
    assert songsdb.find(con, "New Artist", "New Song")["duration_sec"] is None
    second = cross_check([_ev("001", "New Artist - New Song.mp3", rec_out="00:01:00:00")], con)
    # Netz-Treffer bleiben bis zur Bestätigung im Editor ein Vorschlag (db_review), die Dauer kommt aus der EDL
    assert second[0]["status"] == "db_review" and second[0]["duration_sec"] == 60


def test_cross_check_unresolved_without_labelcode(tmp_path):
    con = _db(tmp_path)
    results = cross_check([_ev("001", "Nobody - Ghost Track.mp3")], con,
                          resolve_fn=lambda a, t: {"labelcode": ""})
    assert results[0]["status"] == "unresolved"
    assert songsdb.find(con, "Nobody", "Ghost Track") is None  # nichts geschrieben


def test_export_excel_writes_header_rows_and_total(tmp_path):
    cards = [
        {"artist": "Alanis Morissette", "title": "Thank You", "labelname": "MAVERICK",
         "labelcode": "LC01485", "duration_sec": 6},
        {"artist": "AURORA", "title": "Runaway", "labelname": "Decca",
         "labelcode": "LC00171", "duration_sec": 4},
    ]
    out = tmp_path / "export.xlsx"
    export_excel(cards, out)

    wb = load_workbook(out)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Interpret", "Titel", "Komponist", "Texter", "Verlag", "Label", "LC", "Rechterückruf", "ISRC",
                       "Länge Sek.")
    assert rows[1][:2] == ("Alanis Morissette", "Thank You")
    assert rows[1][9] == 6
    assert rows[2][:2] == ("AURORA", "Runaway")
    assert rows[-1][-1] == 10  # Gesamtdauer = 6 + 4


# --- Merge-Optionen (merge_identical / unify_stems) ---

def test_strip_stem_removes_known_markers_only():
    assert strip_stem("Sunrise (Drums)") == "Sunrise"
    assert strip_stem("Sunrise - Bass 2") == "Sunrise"
    assert strip_stem("Sunrise_VOX") == "Sunrise"
    assert strip_stem("Sunrise (Drums) (Wet)") == "Sunrise"   # mehrfach
    # kein Stem-Wort -> unverändert
    assert strip_stem("Sunrise (Live)") == "Sunrise (Live)"
    assert strip_stem("Track - Reprise") == "Track - Reprise"
    assert strip_stem("Bass") == "Bass"                        # nichts vor dem Marker


def test_merge_cards_off_by_default_returns_unchanged():
    cards = [{"artist": "A", "title": "T", "duration_sec": 3},
             {"artist": "A", "title": "T", "duration_sec": 4}]
    assert merge_cards(cards) == cards   # keine Option -> identisch


def test_merge_identical_sums_duration_and_counts():
    cards = [{"artist": "A", "title": "Song", "labelcode": "LC1", "duration_sec": 3, "status": "found_in_db"},
             {"artist": "A", "title": "Song", "duration_sec": 4, "status": "found_in_db"},
             {"artist": "B", "title": "Other", "duration_sec": 5, "status": "found_in_db"}]
    out = merge_cards(cards, merge_identical=True)
    assert len(out) == 2
    song = next(c for c in out if c["title"] == "Song")
    assert song["duration_sec"] == 7 and song["merged_count"] == 2
    assert song["labelcode"] == "LC1"           # Metadaten aus erster Karte erhalten
    assert next(c for c in out if c["title"] == "Other")["duration_sec"] == 5


def test_unify_stems_merges_stem_variants_to_base():
    cards = [{"artist": "A", "title": "Groove (Drums)", "duration_sec": 3},
             {"artist": "A", "title": "Groove (Bass)", "duration_sec": 4},
             {"artist": "A", "title": "Groove - Vox", "duration_sec": 2}]
    out = merge_cards(cards, unify_stems=True)
    assert len(out) == 1
    assert out[0]["title"] == "Groove" and out[0]["duration_sec"] == 9 and out[0]["merged_count"] == 3


def test_merge_keeps_ungroupable_cards_separate():
    cards = [{"artist": "", "title": "", "clip_name": "weird", "duration_sec": 1, "status": "no_clip_name"},
             {"artist": "A", "title": "T", "duration_sec": 2, "status": "found_in_db"},
             {"artist": "A", "title": "T", "duration_sec": 2, "status": "found_in_db"}]
    out = merge_cards(cards, merge_identical=True)
    assert len(out) == 2   # die nicht-gruppierbare Karte bleibt + eine zusammengefasste
    assert any(c["status"] == "no_clip_name" for c in out)


def test_merge_identical_groups_unparsable_clips_by_filename():
    # Production-Music-Cues ohne Interpret/Titel (nicht als „Artist - Titel" parsebar)
    # werden über den Clip-Namen zusammengefasst — gleiche Datei = eine Zeile, Dauer summiert.
    cards = [
        {"artist": "", "title": "", "clip_name": "ZZQ_TESTLOOP_MUSIK_02", "duration_sec": 1, "status": "unparsable_clip_name"},
        {"artist": "", "title": "", "clip_name": "ZZQ_TESTLOOP_MUSIK_02", "duration_sec": 3, "status": "unparsable_clip_name"},
        {"artist": "", "title": "", "clip_name": "ZZQ_TESTLOOP_MUSIK_02", "duration_sec": 2, "status": "unparsable_clip_name"},
    ]
    out = merge_cards(cards, merge_identical=True)
    assert len(out) == 1
    assert out[0]["duration_sec"] == 6 and out[0]["merged_count"] == 3
    assert out[0]["clip_name"] == "ZZQ_TESTLOOP_MUSIK_02"


def test_merge_unparsable_clips_untouched_without_merge_identical():
    # Ohne merge_identical (nur unify_stems) bleiben unparsebare Clips einzeln —
    # Stem-Vereinheitlichung greift nur auf lesbaren Titeln.
    cards = [
        {"artist": "", "title": "", "clip_name": "X_LOOP", "duration_sec": 1},
        {"artist": "", "title": "", "clip_name": "X_LOOP", "duration_sec": 1},
    ]
    out = merge_cards(cards, unify_stems=True)
    assert len(out) == 2


def test_merge_preserves_order():
    cards = [{"artist": "A", "title": "First", "duration_sec": 1},
             {"artist": "B", "title": "Second", "duration_sec": 1},
             {"artist": "A", "title": "First", "duration_sec": 1}]
    out = merge_cards(cards, merge_identical=True)
    assert [c["title"] for c in out] == ["First", "Second"]


def test_export_csv_has_header_rows_and_total(tmp_path):
    cards = [{"artist": "AURORA", "title": "Runaway", "labelcode": "LC00171", "duration_sec": 4}]
    out = tmp_path / "cue.csv"
    export_csv(cards, out)
    lines = out.read_text(encoding="utf-8-sig").splitlines()
    assert lines[0].startswith("Interpret,")
    assert lines[1].startswith("AURORA,Runaway")
    assert lines[-1].startswith(",,,,,,,,Gesamtdauer:,4")     # 8 Leerzellen vor 'Gesamtdauer:' (10 Spalten)


def test_export_lc_column_is_always_lc_plus_digits(tmp_path):
    """Im Cue-Sheet-Export steht ein Labelcode immer als 'LC' + Ziffern ohne Leerzeichen
    (normalize.to_filename_lc), egal auf welchem Weg er auf die Karte kam: DB-Treffer 'LC01485', Dateiname
    oder Seed '06406', Liste 'LC 24584', Kurzform '171'. Leer bleibt leer, kein Code bleibt sichtbar."""
    from core.normalize import export_lc
    assert [export_lc(v) for v in ("LC01485", "06406", "LC 24584", "171", "LC104994", "", None, "unbekannt")] == \
        ["LC01485", "LC06406", "LC24584", "LC00171", "LC104994", "", "", "unbekannt"]
    cards = [{"artist": "A", "title": "1", "labelcode": "LC01485", "duration_sec": 1},
             {"artist": "B", "title": "2", "labelcode": "06406", "duration_sec": 1},
             {"artist": "C", "title": "3", "labelcode": "LC 24584", "duration_sec": 1},
             {"artist": "D", "title": "4", "labelcode": "", "duration_sec": 1}]
    lc = EXPORT_COLUMNS.index("LC")
    assert [r[lc] for r in _export_rows(cards)[0]] == ["LC01485", "LC06406", "LC24584", ""]
    export_csv(cards, tmp_path / "c.csv")
    import csv
    with open(tmp_path / "c.csv", encoding="utf-8-sig", newline="") as f:
        assert [r[lc] for r in list(csv.reader(f))[1:5]] == ["LC01485", "LC06406", "LC24584", ""]
    export_excel(cards, tmp_path / "c.xlsx")
    ws = load_workbook(tmp_path / "c.xlsx").active
    assert [r[lc] for r in list(ws.iter_rows(values_only=True))[1:5]] == ["LC01485", "LC06406", "LC24584", None]


def test_cross_check_every_card_carries_rechterueckrufe(tmp_path):
    """Jede Cue-Karte trägt `rechterueckrufe` aus der Label-Datenbank zum LC ('' ohne LC oder
    ohne R-Code) — auch DB-Treffer, nicht nur LCs aus dem Dateinamen."""
    con = _db(tmp_path)
    songsdb.upsert(con, interpret="AURORA", titel="Runaway", labelname="Decca", labelcode="LC00171")
    con.commit()
    master = {"00171": ("Decca", "R6"), "00121": ("Atlantic", "R6;R9;R12;R13"), "10192": ("DOMINO", "")}
    res = cross_check([_ev("001", "AURORA - Runaway.wav"),
                       _ev("002", "Zzq Fremd - Song Mit Code - [Atlantic_LC00121].wav"),
                       _ev("003", "Zzq Fremd - Domino Song - [DOMINO_LC10192].wav"),
                       _ev("004", "Nobody - Unknown.wav"), _ev("005", "")], con,
                      label_lookup=lambda c: master.get(c, ("", "")), dedupe=False)
    assert [c["rechterueckrufe"] for c in res] == ["R6", "R6;R9;R12;R13", "", "", ""]


def test_attach_rights_follows_the_current_lc():
    """Nach dem Import kann sich der LC ändern (Zusammenfassen, Editor): der R-Code kommt aus dem aktuellen LC,
    nie von der alten Karte. Ohne Label-Datenbank bleibt ein vorhandener Wert stehen."""
    from core.edl_import import attach_rights
    master = {"00121": ("Atlantic", "R6;R9"), "00171": ("Decca", "R6")}
    cards = [{"labelcode": "LC00171", "rechterueckrufe": "R6;R9"}, {"labelcode": "", "rechterueckrufe": "R6"},
             {"labelcode": "00121"}]
    attach_rights(cards, lambda c: master.get(c, ("", "")))
    assert [c["rechterueckrufe"] for c in cards] == ["R6", "", "R6;R9"]
    keep = [{"labelcode": "00121", "rechterueckrufe": "R6;R9"}, {"labelcode": ""}]
    attach_rights(keep, None)
    assert [c["rechterueckrufe"] for c in keep] == ["R6;R9", ""]


def test_export_has_rechterueckruf_column_after_lc(tmp_path):
    """Spalte 'Rechterückruf' direkt nach der LC-Spalte, in CSV, XLSX und PDF."""
    cards = [{"artist": "Zzq Band", "title": "Nordwand", "labelname": "Atlantic", "labelcode": "00121",
              "rechterueckrufe": "R6;R9;R12;R13", "duration_sec": 5},
             {"artist": "AURORA", "title": "Runaway", "labelcode": "LC00171", "duration_sec": 4}]
    assert EXPORT_COLUMNS.index("Rechterückruf") == EXPORT_COLUMNS.index("LC") + 1
    rr = EXPORT_COLUMNS.index("Rechterückruf")
    rows, total = _export_rows(cards)
    assert [r[rr] for r in rows] == ["R6;R9;R12;R13", ""] and total == 9
    export_excel(cards, tmp_path / "c.xlsx")
    data = list(load_workbook(tmp_path / "c.xlsx").active.iter_rows(values_only=True))
    assert data[0][rr] == "Rechterückruf" and data[1][rr] == "R6;R9;R12;R13" and data[-1][-1] == 9
    export_pdf(cards, tmp_path / "c.pdf")
    import fitz
    text = "".join(fitz.open(tmp_path / "c.pdf")[0].get_text().split())
    assert "Rechterückruf" in text and "R6;R9;R12;R13" in text


def test_export_pdf_is_valid_pdf(tmp_path):
    cards = [{"artist": "AURORA", "title": "Runaway", "labelcode": "LC00171", "duration_sec": 4}]
    out = tmp_path / "cue.pdf"
    export_pdf(cards, out)
    assert out.read_bytes()[:5] == b"%PDF-"
    import fitz
    doc = fitz.open(out)
    assert "Cue-Sheet" in doc[0].get_text() and "AURORA" in doc[0].get_text()
    doc.close()


def test_merge_identical_fills_verlag_and_export_row_carries_it(tmp_path):
    cards = [
        {"artist": "A", "title": "T", "clip_name": "a.wav", "duration_sec": 2, "verlag": ""},
        {"artist": "A", "title": "T", "clip_name": "a.wav", "duration_sec": 3, "verlag": "Edition X"},
    ]
    merged = merge_cards(cards, merge_identical=True)
    assert len(merged) == 1 and merged[0]["verlag"] == "Edition X" and merged[0]["duration_sec"] == 5
    rows, total = _export_rows(merged)
    assert rows[0][EXPORT_COLUMNS.index("Verlag")] == "Edition X" and total == 5


def test_cross_check_catalog_hint_for_library_clip_names(tmp_path):
    from core.resolver import Candidate
    con = _db(tmp_path)
    cat = lambda text: (Candidate(labelcode="10347", label="CMG", confidence=0.5, sources=["GVL-Kürzel CEZ"])
                        if text.startswith("CEZ") else None)
    results = cross_check([_ev("001", "CEZ4030_1.mp3"), _ev("002", "justonewordnoartist.mp3"),
                           _ev("003", "Unknown Artist - Unknown Song.mp3")], con, catalog_fn=cat)
    assert results[0]["status"] == "catalog_hint" and results[0]["labelcode"] == "10347"
    assert results[0]["labelname"] == "CMG" and results[0]["catalog_hint"] == "GVL-Kürzel CEZ"
    assert results[0]["artist"] == "" and results[0]["title"] == "CEZ4030_1"   # nicht in Interpret '1' zerlegt
    assert results[1]["status"] == "unparsable_clip_name" and results[1]["labelcode"] == ""   # kein Kürzel
    assert results[2]["status"] == "missing_no_resolver"                                     # normaler Cue


def test_cross_check_strips_directory_from_source_file(tmp_path):
    # SOURCE FILE trägt bei Premiere/Resolve oft den Pfad; der Ordner darf nicht als
    # Interpret enden ("/audio/AURORA"). Der Clip-Name auf der Karte bleibt unverändert.
    con = _db(tmp_path)
    songsdb.upsert(con, interpret="AURORA", titel="Runaway", labelname="Decca", labelcode="LC00171")
    con.commit()
    for clip in ("/audio/AURORA - Runaway.wav", "C:\\Musik\\AURORA - Runaway.wav"):
        results = cross_check([_ev("001", clip)], con)
        assert results[0]["artist"] == "AURORA", clip
        assert results[0]["status"] == "found_in_db", clip
        assert results[0]["clip_name"] == clip


def test_clip_basename():
    from core.edl_import import clip_basename
    assert clip_basename("/audio/AURORA - Runaway.wav") == "AURORA - Runaway.wav"
    assert clip_basename("C:\\Musik\\CEZ4030_1.mp3") == "CEZ4030_1.mp3"
    assert clip_basename("Thank You.mp3") == "Thank You.mp3"
    assert clip_basename("  ") == ""


def test_cross_check_ambiguous_is_not_written(tmp_path):
    """Margin-Gate im EDL-Fluss: liegen zwei Label gleichauf, wird KEIN Zufallssieger in die
    songs.db geschrieben. Die Karte trägt die Kandidaten, der Mensch wählt im Editor."""
    con = _db(tmp_path)
    tied = [{"labelcode": "00121", "label": "Atlantic", "confidence": 0.5},
            {"labelcode": "07552", "label": "REPUBLIC", "confidence": 0.5}]
    results = cross_check([_ev("001", "Nobody - Two Labels.mp3")], con,
                          resolve_fn=lambda a, t: {"labelcode": "", "ambiguous": True, "tied": tied,
                                                   "komponist": "X", "texter": "Y", "isrc": "USX1"})
    card = results[0]
    assert card["status"] == "ambiguous"
    assert card["labelcode"] == "" and card["tied"] == tied
    assert card["komponist"] == "X" and card["texter"] == "Y" and card["isrc"] == "USX1"   # Autoren kommen mit
    assert songsdb.find(con, "Nobody", "Two Labels") is None                                 # nichts geschrieben


def test_cross_check_ambiguous_keeps_catalog_hint_as_candidate(tmp_path):
    """Gleichstand-Zweig: ein Katalog-Kürzel im Clip-Namen geht als
    drittes Signal in die Kandidatenliste — LC bleibt leer, nichts wird geschrieben."""
    from core.resolver import Candidate
    con = _db(tmp_path)
    cat = lambda text: (Candidate(labelcode="10347", label="CMG", confidence=0.5, sources=["GVL-Kürzel CEZ"])
                        if "CEZ" in text else None)
    tied = [{"labelcode": "00121", "label": "Atlantic", "confidence": 0.5},
            {"labelcode": "07552", "label": "REPUBLIC", "confidence": 0.5}]
    results = cross_check([_ev("001", "Nobody - CEZ4030 Song.mp3")], con, catalog_fn=cat,
                          resolve_fn=lambda a, t: {"labelcode": "", "ambiguous": True, "tied": tied})
    card = results[0]
    assert card["status"] == "ambiguous" and card["labelcode"] == ""
    assert card["catalog_hint"] == "GVL-Kürzel CEZ"
    assert [c["labelcode"] for c in card["tied"]] == ["00121", "07552", "10347"]
    assert songsdb.find(con, "Nobody", "CEZ4030 Song") is None


def test_cross_check_index_hit_fills_card_without_writing(tmp_path):
    """Vorschlag aus dem lokalen Track-Index: Karte gefüllt, Status index_hit, KEIN songs.db-Eintrag."""
    con = _db(tmp_path)
    results = cross_check([_ev("001", "New Artist - New Song.mp3")], con,
                          resolve_fn=lambda a, t: {"labelname": "REPUBLIC", "labelcode": "LC07552",
                                                   "komponist": "X", "texter": "", "index_hit": True})
    assert results[0]["status"] == "index_hit" and results[0]["labelcode"] == "LC07552" and results[0]["komponist"] == "X"
    assert songsdb.find(con, "New Artist", "New Song") is None


def test_cross_check_finds_library_track_by_file_name(tmp_path):
    """Library-Clip „CEZ4030_1“ (nur Katalognummer, kein Interpret im Namen): über library_files sofort die
    verifizierte Zeile, vor Parse-Urteil, Katalog-Hint und Netz (Lieferung Cezame)."""
    con = _db(tmp_path)
    songsdb.import_library_tracks(con, [{"file": "CEZ4030_1.mp3", "interpret": "Pierre-André Athané", "titel": "A Sunny Day",
                                         "komponist": "Pierre-André Athané", "verlag": "Frederic Leibovitz Editeur",
                                         "labelname": "CMG", "labelcode": "LC10347", "album": "Little French Fancies",
                                         "isrc": "FR26F0303350", "duration_sec": 125, "catno": "CEZ4030", "version": "Main version",
                                         "sublibrary": "Cézame Trailers"}],
                                  library="cezame", date_verified="2026-09-20")
    called = []
    results = cross_check([_ev("001", "/audio/CEZ4030_1.wav"), _ev("002", "cez4030_1"), _ev("003", "CEZ4030_9.mp3")], con,
                          resolve_fn=lambda a, t: called.append((a, t)) or None,
                          catalog_fn=lambda text: None, dedupe=False)
    for card in results[:2]:
        assert card["status"] == "found_in_db" and card["matched_by"] == "file"
        assert card["artist"] == "Pierre-André Athané" and card["title"] == "A Sunny Day"
        assert card["labelcode"] == "LC10347" and card["labelname"] == "CMG" and card["komponist"] == "Pierre-André Athané"
        assert card["isrc"] == "FR26F0303350" and card["library"] == "cezame" and card["catno"] == "CEZ4030"
        assert card["sublibrary"] == "Cézame Trailers"
    assert results[2]["status"] == "unparsable_clip_name" and called == []           # unbekannte Nummer: wie bisher


def test_cross_check_library_file_without_lc_still_gets_catalog_proposal(tmp_path):
    """Library-Dateien kommen teils ohne Code (Hooked, Capitol Studio Masters …). Der Lieferungstreffer darf den
    Katalog-Vorschlag nicht verdrängen: erst das Kürzel im Clip-Namen, sonst der Sub-Library-Name der Lieferung
    als Label-Signal. Der LC bleibt Vorschlag (gelb), Titel und Autoren kommen verifiziert aus der Lieferung."""
    from core.resolver import Candidate
    con = _db(tmp_path)
    songsdb.import_library_tracks(con, [
        {"file": "CAP_108_3_Nordwand_Zzq_1373121.mp3", "interpret": "Someone", "titel": "Nordwand",
         "komponist": "Someone", "sublibrary": "Capitol Studio Masters"},
        {"file": "HKD_7_7_Zzq_Goes_North_Testerin_1420031.mp3", "interpret": "Erika Testerin",
         "titel": "Zzq Goes North", "sublibrary": "Hooked"},
        {"file": "XYZ_1_1_Nameless_1999999.mp3", "interpret": "Nobody", "titel": "Nameless", "sublibrary": "Unbekannt"},
    ], library="upm", date_verified="2026-10-08")
    asked = []

    def cat(text, tpub=""):
        asked.append((text, tpub))
        if "CAP108" in text:
            return Candidate(labelcode="92154", label="Capitol Studio Masters", confidence=0.5,
                             sources=["Library-Kürzel CAP (Universal Production Music)"])
        if tpub == "Hooked":
            return Candidate(labelcode="96451", label="Hooked", confidence=0.65, sources=["TPUB = Label"])
        return None

    res = cross_check([_ev("001", "UPM_CAP108_3_NORDWAND_ZZQ_1373121.wav"),
                       _ev("002", "UPM_HKD7_7_ZZQ_GOES_NORTH_1420031.wav"),
                       _ev("003", "XYZ_1_1_Nameless_1999999.wav")], con, catalog_fn=cat, dedupe=False)
    assert [c["matched_by"] for c in res] == ["file", "file", "file"]
    assert (res[0]["status"], res[0]["labelcode"], res[0]["labelname"]) == ("catalog_hint", "92154", "Capitol Studio Masters")
    assert res[0]["title"] == "Nordwand" and res[0]["komponist"] == "Someone"
    assert (res[1]["status"], res[1]["labelcode"], res[1]["catalog_hint"]) == ("catalog_hint", "96451", "TPUB = Label")
    assert (res[2]["status"], res[2]["labelcode"]) == ("found_in_db", "")      # kein Signal: offen wie bisher
    assert asked[0] == ("UPM_CAP108_3_NORDWAND_ZZQ_1373121.wav", "")   # Kürzel zuerst, ohne Namen


def test_base_title_strips_version_markers():
    from core.edl_import import base_title
    assert base_title("NORDWAND (INSTRUMENTAL)") == "NORDWAND"
    assert base_title("NORDLICHT (BONUS TRACK) (INSTRUMENTAL)") == "NORDLICHT"
    assert base_title("ZZQ SONG INSTRUMENTAL GUITAR") == "ZZQ SONG"
    assert base_title("LOVE LOVE LOVE") == "LOVE LOVE LOVE" and base_title("") == ""


def test_siblings_inherit_index_hit_and_keep_filename_lc_note():
    """'NORDWAND' trifft den Index, '(INSTRUMENTAL)' nicht: die Fassung erbt den
    Vorschlag (gelb), der abweichende Dateinamen-LC bleibt als Notiz."""
    from core.edl_import import inherit_from_siblings
    cards = [
        {"event_num": "038", "artist": "ZZQ BAND", "title": "NORDWAND", "labelcode": "00162", "labelname": "COLUMBIA",
         "komponist": "Erika Testerin", "texter": "", "status": "index_hit"},
        {"event_num": "036", "artist": "ZZQ BAND", "title": "NORDWAND (INSTRUMENTAL)", "labelcode": "06667",
         "labelname": "SONY", "komponist": "", "status": "filename_lc"},
        {"event_num": "050", "artist": "OTHER", "title": "SONG (INSTRUMENTAL)", "labelcode": "", "status": "unresolved"},
    ]
    assert inherit_from_siblings(cards) == 1
    c = cards[1]
    assert c["status"] == "sibling_hit" and c["labelcode"] == "00162" and c["labelname"] == "COLUMBIA"
    assert c["komponist"] == "Erika Testerin" and c["sibling_of"] == "038" and "LC06667" in c["note"]
    assert cards[2]["status"] == "unresolved"                                    # kein Geschwister, kein Treffer


def test_siblings_fall_back_to_base_title_lookup():
    from core.edl_import import inherit_from_siblings
    cards = [{"event_num": "1", "artist": "ZZQ BAND", "title": "ZZQ SONG INSTRUMENTAL GUITAR", "labelcode": "", "status": "unresolved"}]
    seen = []
    def local_fn(a, t):
        seen.append((a, t)); return {"labelcode": "02986", "labelname": "DINE-A-MITE JAZZ", "index_hit": True}
    assert inherit_from_siblings(cards, None, local_fn) == 1
    assert seen == [("ZZQ BAND", "ZZQ SONG")] and cards[0]["status"] == "sibling_hit" and cards[0]["sibling_title"] == "ZZQ SONG"


def test_cross_check_picks_candidate_pair_that_hits_locally(tmp_path):
    """'Projekt (SOURCE)_Interpret_Titel_192' — das erste Paar trifft nichts, das zweite den Index."""
    from core import edl, songsdb
    from core.edl_import import cross_check
    con = songsdb.open(tmp_path / "songs.db")
    text = ("TITLE: X\nFCM: NON-DROP FRAME\n001  AX       AA    C        00:00:00:00 00:00:05:00 00:00:00:00 00:00:05:00\n"
            "* FROM CLIP NAME: ZZQ TESTPROJEKT (SOURCE)_ZZQ BAND_NORDWAND LIED_192_L 01846.MP3\n")
    asked = []
    def local_fn(a, t):
        asked.append((a, t))
        return {"labelcode": "01846", "labelname": "UNIVERSAL", "index_hit": True} if (a, t) == ("ZZQ BAND", "NORDWAND LIED") else None
    cards = cross_check(edl.parse_edl(text), con, resolve_fn=local_fn, local_fn=local_fn,
                        label_lookup=lambda c: ("UNIVERSAL", "") if c == "01846" else ("", ""))
    c = cards[0]
    assert (c["artist"], c["title"]) == ("ZZQ BAND", "NORDWAND LIED")
    assert c["status"] == "index_hit" and c["labelcode"] == "01846"
    assert asked[0] != ("ZZQ BAND", "NORDWAND LIED")            # das Projekt-Paar kam zuerst dran


def test_cross_check_filename_lc_is_a_proposal_status(tmp_path):
    from core import edl, songsdb
    from core.edl_import import cross_check
    con = songsdb.open(tmp_path / "songs.db")
    text = ("TITLE: X\nFCM: NON-DROP FRAME\n001  AX       AA    C        00:00:00:00 00:00:05:00 00:00:00:00 00:00:05:00\n"
            "* FROM CLIP NAME: ZZQ BAND -  NORDWAND (INSTRUMENTAL)_LC06667.MP3\n")
    cards = cross_check(edl.parse_edl(text), con, resolve_fn=lambda a, t: None, label_lookup=lambda c: ("SONY", ""))
    assert cards[0]["status"] == "filename_lc" and cards[0]["labelcode"] == "06667"


def test_cross_check_unswaps_pair_when_index_matched_swapped(tmp_path):
    """'ERIKA TESTERIN_NORDLICHT_192' wird als Titel_Interpret gelesen; der Index findet es nur vertauscht
    (swapped=True) — die Karte zeigt danach Interpret ERIKA TESTERIN, Titel NORDLICHT."""
    from core import edl, songsdb
    from core.edl_import import cross_check
    con = songsdb.open(tmp_path / "songs.db")
    text = ("TITLE: X\nFCM: NON-DROP FRAME\n001  AX       AA    C        00:00:00:00 00:00:05:00 00:00:00:00 00:00:05:00\n"
            "* FROM CLIP NAME: ERIKA TESTERIN_NORDLICHT_192_LC07005.MP3\n")
    def local_fn(a, t):
        if (a, t) == ("NORDLICHT", "ERIKA TESTERIN"):
            return {"labelcode": "58564", "labelname": "Dumont Dumont", "index_hit": True, "swapped": True}
        return None
    cards = cross_check(edl.parse_edl(text), con, resolve_fn=local_fn, local_fn=local_fn, label_lookup=lambda c: ("", ""))
    assert (cards[0]["artist"], cards[0]["title"]) == ("ERIKA TESTERIN", "NORDLICHT") and cards[0]["labelcode"] == "58564"


def test_pdf_export_schneidet_lange_dateinamen_nicht_ab(tmp_path):
    """Ein Dateiname ohne Leerzeichen als Titel nahm ohne feste Spaltenbreiten die halbe Seite und
    drückte „Länge Sek." aus dem Blatt, bei 113 Zeichen fehlte die Spalte ganz."""
    import fitz
    long = ("ZZQ_TESTPROJEKT_EINS_SOURCE_ZZQ_BAND_NORDWAND_LIED_192_LC01846_"
            "EXTRA_LANGER_DATEINAME_OHNE_LEERZEICHEN")
    cards = [{"artist": "", "title": long, "duration_sec": 5},
             {"artist": "Matt Black", "title": "Scars (A)", "composer": "Matt Black",
              "labelname": "KPM MUSIC", "labelcode": "LC02061", "duration_sec": 6}]
    pdf = tmp_path / "cue.pdf"
    export_pdf(cards, pdf)
    page = fitz.open(pdf)[0]
    words = page.get_text("words")
    assert any(w[4] == "Sek." for w in words), "letzte Spalte fehlt"
    assert any(w[4] == "Gesamtdauer" for w in words)
    assert max(w[2] for w in words) <= page.rect.x1 - 36 + 0.5, "Text läuft über den Rand"
    # Der Name steht vollständig da, nur über mehrere Zeilen umbrochen: die Umbruchhilfen sind
    # unsichtbar, die Textextraktion setzt an jeder Umbruchstelle einen Zeilenwechsel.
    assert long in "".join(page.get_text().split())


def test_pdf_export_haelt_spaltenbreiten_auch_ohne_trennzeichen(tmp_path):
    """Die PyMuPDF-Story ignoriert Prozentbreiten und setzt jede Spalte mindestens so breit wie ihr längstes
    unumbrechbares Stück. Ein Wort ohne _ + - . / (113 oder 200 Zeichen, CamelCase, kommagetrennte Namen) schob
    'ISRC', 'Länge Sek.' und die Gesamtdauer aus dem Blatt. Jetzt: feste Breiten in pt und Umbruchstellen, wo
    ein Stück breiter ist als seine Spalte. Jede Spalte steht dort, wo die Breitentabelle sie hinlegt."""
    import fitz
    from core.edl_import import _PDF_WIDTHS
    camel = "".join(w.capitalize() for w in ("zzq testprojekt eins source zzq band nordwand lied "
                                             "extra langer dateiname ohne leerzeichen").split())
    komma = ",".join(["Johnson", "Smith", "Williams", "Brown", "Jones", "Garcia", "Miller", "Davis", "Rodriguez",
                      "Martinez", "Hernandez", "Lopez", "Gonzalez", "Wilson", "Anderson"])
    base = {"artist": "Matt Black", "title": "Scars (A)", "komponist": "Matt Black", "labelname": "KPM MUSIC",
            "labelcode": "LC02061", "rechterueckrufe": "R6;R9;R12;R13", "isrc": "GBXXX0000001", "duration_sec": 6}
    for extra in ({"title": "x" * 113}, {"title": "x" * 200}, {"title": camel}, {"komponist": komma},
                  {"verlag": "Äöüßéèàçñ" * 6}):
        pdf = tmp_path / "cue.pdf"
        export_pdf([{**base, **extra}], pdf)
        page = fitz.open(pdf)[0]
        words = page.get_text("words")
        limit = page.rect.x1 - 36 + 0.5
        assert max(w[2] for w in words) <= limit, extra
        assert max(d["rect"].x1 for d in page.get_drawings()) <= limit, extra      # Tabellenrahmen im Blatt
        heads = {w[4]: w[0] for w in words if w[1] < 90}
        for h in ("Interpret", "Titel", "Komponist", "Label", "LC", "Rechterückruf", "ISRC", "Länge", "Sek."):
            assert h in heads, (h, extra)
        left = 36.0
        for col, pct in zip(EXPORT_COLUMNS, _PDF_WIDTHS):                              # Spalte an ihrem Platz
            if col.split()[0] in heads:
                assert abs(heads[col.split()[0]] - (left + 5.5)) < 1.5, (col, extra)
            left += pct * (page.rect.width - 72) / 100
        assert any(w[4] == "Gesamtdauer" for w in words)
        text = "".join(page.get_text().split())
        assert (extra.get("title") or extra.get("komponist") or extra.get("verlag")) in text    # nichts abgeschnitten


def test_soft_break_nur_in_langen_wortketten():
    assert _soft_break("Scars (A)") == "Scars (A)"          # kurz: unverändert
    assert _soft_break("Agnes Obel – Dorian") == "Agnes Obel – Dorian"
    out = _soft_break("ZZQ_TESTPROJEKT_EINS+ZZQ+NORDWAND+SERIE_1920X1080")
    assert "​" in out and out.replace("​", "") == "ZZQ_TESTPROJEKT_EINS+ZZQ+NORDWAND+SERIE_1920X1080"


def test_cross_check_split_edit_is_one_cue(tmp_path):
    con = _db(tmp_path)
    evs = [_ev("001", "X - Y.wav", "00:00:00:00", "00:00:10:00"), _ev("002", "X - Y.wav", "00:00:10:00", "00:00:25:00")]
    res = cross_check(evs, con)
    assert len(res) == 1 and res[0]["duration_sec"] == 25 and res[0]["record_in"] == "00:00:00:00"
    assert len(cross_check(evs, con, merge_split_edits=False)) == 2


def test_online_hit_is_written_for_review_and_stays_yellow_on_reimport(tmp_path):
    """Ein Netz-Treffer ist ein Vorschlag: beim nächsten Import 'db_review' (gelb), nicht 'found_in_db'."""
    con = _db(tmp_path)
    resolve = lambda a, t: {"labelcode": "LC00121", "labelname": "Atlantic"}
    first = cross_check([_ev("001", "New Artist - New Song.mp3")], con, resolve_fn=resolve)
    assert first[0]["status"] == "resolved_and_written"
    second = cross_check([_ev("001", "New Artist - New Song.mp3")], con)
    assert second[0]["status"] == "db_review"
    # Im Editor gespeichert -> verifiziert
    row = songsdb.find_row(con, "New Artist", "New Song")
    songsdb.update_song(con, row["id"], {"labelname": "Atlantic"})
    assert cross_check([_ev("001", "New Artist - New Song.mp3")], con)[0]["status"] == "found_in_db"


def test_isrc_in_clip_name_is_used_and_not_read_as_title(tmp_path):
    con = _db(tmp_path)
    songsdb.upsert(con, interpret="ABBA", titel="Dancing Queen", labelname="Polar", labelcode="LC07032",
                   isrc="SE-AYD-76-01020")
    con.commit()
    hit = cross_check([_ev("001", "SEAYD7601020_dq_final.wav")], con)[0]
    assert hit["status"] == "found_in_db" and hit["matched_by"] == "isrc" and hit["title"] == "Dancing Queen"
    miss = cross_check([_ev("002", "GBAYE7700001_Unknown.wav")], con)[0]
    assert miss["isrc"] == "GBAYE7700001"
    assert "GBAYE7700001" not in (miss["title"] + miss["artist"])
